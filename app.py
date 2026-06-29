"""
Inventario PYME - Backend Flask
Requisitos: pip install flask flask-cors mysql-connector-python openpyxl
Uso: python app.py
"""

from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
import mysql.connector
import openpyxl
import os
import random

app = Flask(__name__, static_folder=".")
CORS(app)

# ── Configuración de base de datos (Dinámica) ─────────────────

global_db_config = {
    "host": os.getenv("DB_HOST", "localhost"),
    "user": os.getenv("DB_USER","root"),
    "password": os.getenv("DB_PASSWORD", ""),
    "database": os.getenv("DB_NAME", "pyme_db")
}

def get_db():
    """Retorna una nueva conexión a MySQL usando datos globales."""
    try:
        return mysql.connector.connect(**global_db_config)
    except mysql.connector.Error:
        return None

@app.route("/api/status", methods=["GET"])
def connection_status():
    conn = get_db()
    if conn is not None:
        conn.close()
        return jsonify({"connected": True})
    return jsonify({"connected": False}), 200


# ── Rutas de la interfaz ────────────────────────────────────────
@app.route("/")
def index():
    return send_from_directory(".", "index.html")


# ── API Universales: Descubrimiento de Esquema ──────────────────
@app.route("/api/schema", methods=["GET"])
def get_schema():
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    cur = conn.cursor(dictionary=True)
    
    try:
        cur.execute("SHOW TABLES")
        tables_res = cur.fetchall()
        tables = [list(row.values())[0] for row in tables_res]
        
        schema = {}
        for table in tables:
            cur.execute(f"SHOW COLUMNS FROM `{table}`")
            columns = cur.fetchall()
            schema[table] = columns
            
        cur.close()
        conn.close()
        return jsonify(schema)
    except mysql.connector.Error as err:
        return jsonify({"error": str(err)}), 400


# ── API: Leer datos dinámicos ───────────────────────────────────
@app.route("/api/datos/<tabla>", methods=["GET"])
def get_datos(tabla):
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    
    limite = int(request.args.get("limit", 100))
    offset = int(request.args.get("offset", 0))

    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(f"SELECT * FROM `{tabla}` LIMIT %s OFFSET %s", (limite, offset))
        datos = cur.fetchall()
        cur.close()
        conn.close()
        return jsonify(datos)
    except mysql.connector.Error as err:
        return jsonify({"error": str(err)}), 400


# ── API: Insertar dato dinámico ─────────────────────────────────
@app.route("/api/datos/<tabla>", methods=["POST"])
def pos_datos(tabla):
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    
    data = request.get_json()
    if not data: return jsonify({"error": "No payload"}), 400
    
    keys = list(data.keys())
    values = tuple(data.values())
    placeholders = ", ".join(["%s"] * len(values))
    cols = ", ".join([f"`{k}`" for k in keys])
    
    cur = conn.cursor()
    try:
        query = f"INSERT INTO `{tabla}` ({cols}) VALUES ({placeholders})"
        cur.execute(query, values)
        conn.commit()
        new_id = cur.lastrowid
        cur.close()
        conn.close()
        return jsonify({"mensaje": "Insertado correctamente", "id": new_id}), 201
    except mysql.connector.Error as err:
        return jsonify({"error": str(err)}), 400


# ── API: Actualizar dato dinámico ───────────────────────────────
@app.route("/api/datos/<tabla>/<pk_col>/<pk_val>", methods=["PUT"])
def update_dato(tabla, pk_col, pk_val):
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    
    data = request.get_json()
    if not data: return jsonify({"error": "No payload"}), 400
    
    set_clause = ", ".join([f"`{k}`=%s" for k in data.keys()])
    values = list(data.values())
    values.append(pk_val)
    
    cur = conn.cursor()
    try:
        query = f"UPDATE `{tabla}` SET {set_clause} WHERE `{pk_col}` = %s"
        cur.execute(query, values)
        conn.commit()
        afectadas = cur.rowcount
        cur.close()
        conn.close()
        return jsonify({"mensaje": f"{afectadas} fila(s) actualizada(s)"})
    except mysql.connector.Error as err:
        return jsonify({"error": str(err)}), 400


# ── API: Eliminar dato dinámico ─────────────────────────────────
@app.route("/api/datos/<tabla>/<pk_col>/<pk_val>", methods=["DELETE"])
def delete_dato(tabla, pk_col, pk_val):
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    
    cur = conn.cursor()
    try:
        cur.execute(f"DELETE FROM `{tabla}` WHERE `{pk_col}` = %s", (pk_val,))
        conn.commit()
        afectadas = cur.rowcount
        cur.close()
        conn.close()
        return jsonify({"mensaje": f"{afectadas} fila(s) eliminada(s)"})
    except mysql.connector.Error as err:
        return jsonify({"error": str(err)}), 400


# ── API: Importador Mapeado Dinámico ─────────────────────────────
@app.route("/api/importar_mapeado/<tabla>", methods=["POST"])
def importar_mapeado(tabla):
    conn = get_db()
    if not conn: return jsonify({"error": "No db"}), 500
    
    data = request.get_json()
    if not isinstance(data, list) or len(data) == 0:
         return jsonify({"error": "Se esperaba una lista JSON"}), 400
    
    agregados = 0
    errores = []
    
    cur = conn.cursor()
    keys = list(data[0].keys())
    cols = ", ".join([f"`{k}`" for k in keys])
    placeholders = ", ".join(["%s"] * len(keys))
    query = f"INSERT INTO `{tabla}` ({cols}) VALUES ({placeholders})"
    
    for i, row in enumerate(data):
        try:
            values = tuple(row.get(k) for k in keys)
            cur.execute(query, values)
            agregados += 1
        except mysql.connector.Error as e:
            errores.append(f"Fila {i+1}: {str(e)}")
            
    conn.commit()
    cur.close()
    conn.close()
    
    return jsonify({
        "agregados": agregados,
        "errores": errores,
        "mensaje": f"{agregados} importados."
    })

if __name__ == "__main__":
    port = int(os.getenv("PORT", 5000))
    print(f"Servidor corriendo en http://localhost:{port}")
    app.run(host="0.0.0.0", port=port, debug=False)